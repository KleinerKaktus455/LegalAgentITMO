"""Fill the ITMO scientific-practice template with English project slides."""

from __future__ import annotations

import shutil
from pathlib import Path

from copy import deepcopy

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu, Pt

ROOT = Path(__file__).resolve().parent
SRC = ROOT / "template_slides" / "template.pptx"
DST = ROOT / "ITMO_practice_presentation.pptx"
SHOTS = ROOT / "screenshots"


def set_runs(shape, lines: list[str]) -> None:
    tf = shape.text_frame
    paras = list(tf.paragraphs)
    for i, text in enumerate(lines):
        if i < len(paras):
            p = paras[i]
            if p.runs:
                p.runs[0].text = text
                for run in p.runs[1:]:
                    run.text = ""
            else:
                p.text = text
        else:
            p = tf.add_paragraph()
            p.text = text
            if paras and paras[0].runs:
                src = paras[0].runs[0]
                run = p.runs[0] if p.runs else p.add_run()
                run.text = text
                if src.font.size:
                    run.font.size = src.font.size
                if src.font.name:
                    run.font.name = src.font.name
    for j in range(len(lines), len(paras)):
        for run in paras[j].runs:
            run.text = ""


def set_body(shape, blocks: list[str]) -> None:
    """Replace body text, keeping the first run's font on paragraph 0."""
    tf = shape.text_frame
    first = tf.paragraphs[0]
    font_size = first.runs[0].font.size if first.runs else None
    font_name = first.runs[0].font.name if first.runs else None

    # wipe existing
    for p in tf.paragraphs:
        for run in p.runs:
            run.text = ""

    texts = [b for b in blocks if b is not None]
    for i, text in enumerate(texts):
        p = tf.paragraphs[i] if i < len(tf.paragraphs) else tf.add_paragraph()
        if p.runs:
            p.runs[0].text = text
            for run in p.runs[1:]:
                run.text = ""
        else:
            run = p.add_run()
            run.text = text
            if font_size:
                run.font.size = font_size
            if font_name:
                run.font.name = font_name


def set_table(shape, rows: list[list[str]]) -> None:
    tbl = shape.table
    for r_i, row in enumerate(rows):
        for c_i, val in enumerate(row):
            cell = tbl.cell(r_i, c_i)
            # keep first para/run
            tf = cell.text_frame
            if tf.paragraphs and tf.paragraphs[0].runs:
                tf.paragraphs[0].runs[0].text = val
                for run in tf.paragraphs[0].runs[1:]:
                    run.text = ""
                for extra in tf.paragraphs[1:]:
                    for run in extra.runs:
                        run.text = ""
            else:
                cell.text = val


def set_layout_text(layout, shape_idx: int, lines: list[str]) -> None:
    sh = layout.shapes[shape_idx]
    if sh.has_text_frame:
        set_runs(sh, lines)


def _insert_slide(prs: Presentation, index: int, layout):
    slide = prs.slides.add_slide(layout)
    sld_id = prs.slides._sldIdLst[-1]
    prs.slides._sldIdLst.remove(sld_id)
    prs.slides._sldIdLst.insert(index, sld_id)
    return slide


def _crop_png(src: Path, dst: Path, bottom_drop: float = 0.0) -> Path:
    im = Image.open(src)
    w, h = im.size
    box = (0, 0, w, int(h * (1.0 - bottom_drop)))
    im.crop(box).save(dst)
    return dst


def _add_demo_slide(prs: Presentation, index: int, layout, title: str, caption: str, image: Path) -> None:
    slide = _insert_slide(prs, index, layout)
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        if sh.has_text_frame and "‹#›" in sh.text_frame.text:
            continue
        name = (sh.name or "").lower()
        if "title" in name or sh.top < 1_100_000:
            set_runs(sh, [title])
            break

    # Content box under the title, above the footer / caption.
    box_l, box_t = 360_000, 1_180_000
    box_w, box_h = 11_400_000, 4_850_000
    with Image.open(image) as im:
        iw, ih = im.size
    ratio = iw / ih
    pic_w = box_w
    pic_h = int(pic_w / ratio)
    if pic_h > box_h:
        pic_h = box_h
        pic_w = int(pic_h * ratio)
    pic_l = box_l + (box_w - pic_w) // 2
    slide.shapes.add_picture(str(image), Emu(pic_l), Emu(box_t), Emu(pic_w), Emu(pic_h))

    cap = slide.shapes.add_textbox(Emu(box_l), Emu(6_180_000), Emu(8_400_000), Emu(380_000))
    p = cap.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = caption
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

    for donor in prs.slides:
        for sh in donor.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip() == "‹#›":
                slide.shapes._spTree.append(deepcopy(sh._element))
                return


def main() -> None:
    shutil.copy2(SRC, DST)
    prs = Presentation(str(DST))

    # Thank-you layout
    for layout in prs.slide_layouts:
        if layout.name == "Финальный слайд":
            for sh in layout.shapes:
                if sh.has_text_frame and ("Спасибо" in sh.text_frame.text or "Thank you" in sh.text_frame.text):
                    set_runs(sh, ["Thank you", "for your attention!"])
                    for p in sh.text_frame.paragraphs:
                        for run in p.runs:
                            if run.text.strip():
                                run.font.color.rgb = RGBColor(255, 255, 255)

    s = list(prs.slides)

    # 1 title
    set_runs(s[0].shapes[0], ["Multi-Agent Court-Act Assistant"])
    set_runs(
        s[0].shapes[1],
        [
            "Student:",
            "Anastasia Senichenkova, 01.04.02 Applied Mathematics and Informatics",
            "Scientific supervisor:",
            "to be specified",
            "Scientific consultant:",
            "to be specified",
            "",
        ],
    )

    # 2 problem
    set_runs(s[1].shapes[0], ["Problem statement"])
    set_body(
        s[1].shapes[2],
        [
            "Open corpora of Russian courts of general jurisdiction (SOY) are large, mixed, and hard to search for analogous cases.",
            "A case card stores metadata and sometimes HTML; a parallel corpus splits the act into header, facts (fabula), and operative part. There is no gold outcome label.",
            "Legal LLMs often invent citations. Predicting a judgment from this data would be scientifically weak and practically misleading.",
            "Scientific track: normalize the card into a Case Graph, retrieve comparable practice from an index, and draft a position that may cite only indexed case IDs.",
        ],
    )

    # 3 goal
    set_runs(s[2].shapes[0], ["Goal and tasks"])
    set_body(
        s[2].shapes[2],
        [
            "Goal: build a working multi-agent prototype that analyzes a SOY case card, retrieves analogous practice, and drafts a grounded legal position.",
            "Tasks:",
            "Analyze the JSONL cards and parquet act-splits; define the Case Graph schema.",
            "Implement ETL into SQLite so agents never scan raw gigabytes.",
            "Implement a LangGraph pipeline with playbooks (dossier / analogs / risk / appeal / draft / full).",
            "Evaluate routing, analog precision@5, and a citation Guard; ship CLI and Streamlit UI.",
        ],
    )

    # 4 related work
    set_runs(s[3].shapes[0], ["Related work"])
    set_body(
        s[3].shapes[2],
        [
            "Zhong et al.; Chalkidis et al., LexGLUE / LEGAL-BERT: classification and judgment prediction. Gap: our corpus has no reliable outcome label; a “win probability” would overclaim.",
            "COLIEE; Shao et al., BERT-PLI: learned case retrieval. Gap: we need an explainable ranker on Russian SOY acts, not an English encoder.",
            "Lewis et al., RAG; Cui et al., ChatLaw: generate from retrieved text. Gap: generation must not invent case IDs.",
            "Wu et al., AutoGen; LangGraph: LLM multi-agent stacks. We use an explicit linear graph: nodes outside the plan skip.",
            "Magesh et al.: leading legal AI tools hallucinate citations. This motivates Guard against the SQLite index.",
            "Takeaway: the gap is grounded retrieval and drafting — not another outcome classifier.",
        ],
    )

    # 5 methods
    set_runs(s[4].shapes[0], ["Methods and tools"])
    set_body(
        s[4].shapes[2],
        [
            "Methods: rule-based legal NLP, Case Graph normalization, lexical information retrieval, multi-agent orchestration with playbooks.",
            "Analog score (explainable): 5·same article + 3·same judge + 2·same region + 1.5·same instance + Jaccard overlap of fabula tokens.",
            "Tools: Python 3.11+, LangGraph, Pydantic, PyArrow, SQLite (WAL, batch upsert), Streamlit; optional Ollama (OpenAI-compatible API).",
            "Data: docs.json SOY cards; correct_df_splitted_text.parquet (200 parts, text_1 / text_2 / text_3).",
        ],
    )

    # 6 novelty
    set_runs(s[5].shapes[0], ["Scientific novelty"])
    set_body(
        s[5].shapes[2],
        [
            "The task is restated as Case Graph construction and analog retrieval, not judgment prediction.",
            "A linear LangGraph with six playbooks: unused nodes skip, so one graph serves dossier, practice search, risk, appeal, and drafting.",
            "Hybrid analog score is transparent (article / judge / region / instance / Jaccard) and runs without embeddings.",
            "Guard accepts a draft only if cited case IDs exist in the index — a direct response to citation hallucination.",
            "The pipeline is usable without an LLM key (rules + index statistics); the model is optional wording, not the source of facts.",
        ],
    )

    # 7 solution 1
    set_runs(s[6].shapes[0], ["Solution: data and index"])
    set_body(
        s[6].shapes[2],
        [
            "JSONL cards and parquet rows are normalized into one Case Graph (id, court, judge, region, charges, timeline, fabula, operative part).",
            "ingest writes SQLite: sources parquet / docs / both, limit, reset; existing parquet IDs are skipped; batch insert + WAL.",
            "Full load: 1,862,811 cases from 200 parquet parts. Agents query case_store.sqlite only.",
            "Corpus mix: civil 1,609,786; administrative 161,766; unknown 91,121; criminal 138 — so the system is not a criminal-outcome predictor.",
        ],
    )

    # 8 solution 2
    set_runs(s[7].shapes[0], ["Solution: agents and playbooks"])
    set_body(
        s[7].shapes[2],
        [
            "Linear graph: Orchestrator → Parser → Timeline → Act Reader → Qualification → Retriever → Risk → Appeal → Draft → Guard → Respond.",
            "Playbooks: dossier (card only); analogs; risk; appeal; draft; full. Orchestrator picks by query rules; an explicit --playbook is never overridden.",
            "Retriever ranks by the hybrid score; Respond prints markdown only from AgentState; Guard checks cited IDs.",
            "Interfaces share the same functions: CLI (ingest / status / ask) and Streamlit UI.",
        ],
    )

    # 9 results
    set_runs(s[8].shapes[0], ["Main results"])
    set_body(
        s[8].shapes[2],
        [
            "Working artifact: legal_multiagent (CLI, Streamlit, README). Full index: 1.86M cases (~69 GB SQLite).",
            "30 automated tests passed (normalization, retrieval, routing, Guard, Risk). No invented facts in Respond.",
            "Retriever: ≥ 60% of top-5 analogs share article and region. Guard rejects unknown case IDs.",
            "Code and README: repository root (cognitive-funnel README; not a public license).",
        ],
    )
    set_table(
        s[8].shapes[3],
        [
            ["Kind", "Records", "Share", "Role in system"],
            ["Civil", "1,609,786", "86.4%", "main analog pool"],
            ["Administrative / other / criminal", "253,025", "13.6%", "sparse; 138 criminal"],
        ],
    )

    # 10 results extra
    set_runs(s[9].shapes[0], ["Main results"])
    set_body(
        s[9].shapes[2],
        [
            "Qualitative checks",
            "Routing: five query types map to the intended playbook.",
            "Respond: final answer contains only numbers and IDs present in state.",
            "Risk: wording states this is not a court forecast; warns when analogs are few.",
            "Limits of the current evaluation",
            "No lawyer-labeled “this analog is useful” set on the full corpus (precision@5 is a proxy).",
            "LLM branches (act reader / drafting) were not run e2e without an API key.",
            "Lexical Jaccard can promote a different article if token overlap is high.",
        ],
    )

    # 11 expected
    set_runs(s[10].shapes[0], ["Expected next results"])
    set_body(
        s[10].shapes[2],
        [
            "Fabula embeddings for analog search, compared with the current hybrid score.",
            "A small lawyer-labeled relevance slice and precision/recall on that slice.",
            "Mocks for Ollama so the full LangGraph path runs in CI.",
            "Submit the Elsevier-format draft (paper/main.tex) to a student / legal-IR workshop.",
        ],
    )

    # 12 conclusion
    set_runs(s[11].shapes[0], ["Conclusion"])
    set_body(
        s[11].shapes[2],
        [
            "The goal is met: a prototype that builds a Case Graph, retrieves analogs from SQLite, and drafts a position without citing missing IDs.",
            "Task 1–2: schema + ETL for 1.86M parquet acts. Task 3: LangGraph playbooks, CLI, UI. Task 4: 30 tests; routing and Guard hold.",
            "The system is an assistant for practice search, not a judgment classifier and not legal advice.",
            "Next: embeddings, human relevance labels, and a workshop submission.",
        ],
    )

    # 13 publication
    set_runs(s[12].shapes[0], ["Publication readiness"])
    set_body(
        s[12].shapes[2],
        [
            "A full manuscript is prepared in Elsevier elsarticle (paper/main.tex + references.bib).",
            "Status: draft / preprint in the repository — not submitted, no acceptance yet.",
            "Venue: to be selected (legal IR, multi-agent systems, or an ITMO student track) within three months.",
            "No reviewer reports yet. The Overleaf-ready source can be compiled as pdfLaTeX.",
        ],
    )

    empty = next(lo for lo in prs.slide_layouts if lo.name == "Пустой слайд")
    query_img = _crop_png(SHOTS / "01_ui_query.png", SHOTS / "01_ui_query_crop.png", bottom_drop=0.30)
    # After “Solution: agents” (index 7) — before Main results.
    _add_demo_slide(
        prs,
        8,
        empty,
        "Demo: Streamlit query",
        "Index of 1,862,811 cases. Query: tax-arrears analog search. Playbook: analogs.",
        query_img,
    )
    _add_demo_slide(
        prs,
        9,
        empty,
        "Demo: analog table and graph steps",
        "20 ranked analogs. Unused LangGraph nodes skip (parser → … → Guard).",
        SHOTS / "03_ui_steps.png",
    )

    prs.save(str(DST))
    print(f"wrote {DST}")


if __name__ == "__main__":
    main()
